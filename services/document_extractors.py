from __future__ import annotations

import csv
import io
from typing import Optional

def extract_text(filename: str, content_type: Optional[str], data: bytes) -> str:
    name = (filename or "").lower()
    ctype = (content_type or "").lower()

    # TXT / MD
    if name.endswith((".txt", ".md")) or "text/plain" in ctype:
        return data.decode("utf-8", errors="ignore")

    # CSV
    if name.endswith(".csv") or "text/csv" in ctype:
        return _extract_csv(data)

    # PDF
    if name.endswith(".pdf") or "application/pdf" in ctype:
        return _extract_pdf(data)

    # DOCX
    if name.endswith(".docx") or "application/vnd.openxmlformats-officedocument.wordprocessingml.document" in ctype:
        return _extract_docx(data)

    # XLSX
    if name.endswith(".xlsx") or "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet" in ctype:
        return _extract_xlsx(data)

    # fallback best-effort
    return data.decode("utf-8", errors="ignore")


def _extract_pdf(data: bytes) -> str:
    try:
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(data))
        out = []
        for page in reader.pages:
            out.append(page.extract_text() or "")
        return "\n".join(out).strip()
    except Exception:
        return ""


def _extract_docx(data: bytes) -> str:
    try:
        from docx import Document
        doc = Document(io.BytesIO(data))
        return "\n".join([p.text for p in doc.paragraphs if p.text]).strip()
    except Exception:
        return ""


def _extract_xlsx(data: bytes) -> str:
    try:
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
        out = []
        for ws in wb.worksheets:
            out.append(f"--- Sheet: {ws.title} ---")
            for row in ws.iter_rows(values_only=True):
                line = " | ".join("" if v is None else str(v) for v in row)
                if line.strip():
                    out.append(line)
        return "\n".join(out).strip()
    except Exception:
        return ""


def _extract_csv(data: bytes) -> str:
    try:
        text = data.decode("utf-8", errors="ignore")
        f = io.StringIO(text)
        reader = csv.reader(f)
        out = []
        for row in reader:
            out.append(" | ".join(row))
        return "\n".join(out).strip()
    except Exception:
        return ""
